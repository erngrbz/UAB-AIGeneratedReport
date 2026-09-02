from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from typing import List, Optional
import requests
from bs4 import BeautifulSoup
import json
import urllib.request
import os
import base64

from weasyprint import HTML
from docx import Document
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE

app = FastAPI(title="Haber Analiz ve Raporlama API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

os.makedirs("reports", exist_ok=True)
app.mount("/reports", StaticFiles(directory="reports"), name="reports")

URL = "http://localhost:8000/v1/chat/completions"
MODEL_NAME = "qwen3.6-27b-nvfp4-dflash10"
BASE_URL = "http://localhost:8080/api/accident"
BASE_URL2 = "http://localhost:8081/api/accident"

app_data = {
    "bus_license_plates": [],
    "bus_count": 0,
    "news_date": "",
    "news_date2": "",
    "news_date3": "",
    "summary_sentences": []
}

class HaberRequest(BaseModel):
    url: str

class AracBilgisi(BaseModel):
    plate: str
    accident_date: Optional[str]

class RaporRequest(BaseModel):
    buses: List[AracBilgisi]

def generate_pdf_report(output_filename, plate, date, text, logo_path):
    logo_base64 = ""
    if os.path.exists(logo_path):
        with open(logo_path, "rb") as img_file:
            logo_base64 = base64.b64encode(img_file.read()).decode('utf-8')
            
    paragraphs = text.split('\n')
    formatted_html_paragraphs = ""
    for para_text in paragraphs:
        para_text = para_text.strip()
        if not para_text:
            continue
        formatted_html_paragraphs += f"<p>{para_text}</p>"

    html_template = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            @page {{
                size: A4;
                margin: 15mm;
                background-color: #ffffff;
            }}
            body {{
                font-family: 'Times New Roman', Times, serif;
                font-size: 11.5pt;
                color: #000;
                line-height: 1.45;
                margin: 0;
                padding: 0;
            }}
            table.report-table {{
                width: 100%;
                border-collapse: collapse;
                table-layout: fixed;
            }}
            table.report-table, .report-table th, .report-table td {{
                border: 1px solid #000000;
            }}
            .report-table td {{
                padding: 8px 10px;
                vertical-align: middle;
            }}
            .logo-cell {{
                width: 28%;
                text-align: center;
                vertical-align: middle;
                padding: 0px !important;
            }}
            .title-cell {{
                width: 72%;
                text-align: center;
                vertical-align: middle;
            }}
            .logo-img {{
                max-width: 100%;
                height: auto;
                display: block;
                margin: 0 auto;
            }}
            .main-title {{
                font-family: 'Times New Roman', Times, serif;
                font-size: 18pt;
                font-weight: bold;
                color: #C92228;
                letter-spacing: 0.5px;
            }}
            .label-cell {{
                width: 28%;
                font-weight: bold;
            }}
            .value-cell {{
                width: 72%;
            }}
            .section-header {{
                text-align: center;
                font-weight: bold;
                font-size: 12pt;
                background-color: #f2f2f2;
            }}
            .content-cell {{
                padding: 15px !important;
                text-align: justify;
                vertical-align: top;
            }}
            .content-cell p {{
                text-indent: 7mm;
                margin-top: 0;
                margin-bottom: 12pt;
            }}
        </style>
    </head>
    <body>
        <table class="report-table">
            <tr>
                <td class="logo-cell">
                    {"<img src='data:image/png;base64," + logo_base64 + "' class='logo-img'>" if logo_base64 else "[Logo Bulunamadı]"}
                </td>
                <td class="title-cell">
                    <div class="main-title">BİLGİ NOTU</div>
                </td>
            </tr>
            <tr>
                <td class="label-cell">Konu</td>
                <td class="value-cell">Otobüs Kazası – {plate} Plakalı Taşıt</td>
            </tr>
            <tr>
                <td class="label-cell">Hazırlanış Tarihi</td>
                <td class="value-cell">{date}</td>
            </tr>
            <tr>
                <td colspan="2" class="section-header">AÇIKLAMA</td>
            </tr>
            <tr>
                <td colspan="2" class="content-cell">
                    {formatted_html_paragraphs}
                </td>
            </tr>
        </table>
    </body>
    </html>
    """
    HTML(string=html_template).write_pdf(output_filename)

def generate_formal_report_docx(output_filename, plate, date, text, logo_path):
    document = Document()
    
    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(15)
        section.right_margin = Mm(15)
        section.top_margin = Mm(15)
        section.bottom_margin = Mm(15)

    style = document.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11.5)

    topic = f"Otobüs Kazası – {plate} Plakalı Taşıt"

    table = document.add_table(rows=5, cols=2)
    table.style = 'Table Grid'
    
    table.autofit = False
    table.allow_autofit = False
    for cell in table.columns[0].cells:
        cell.width = Mm(57.6)
    for cell in table.columns[1].cells:
        cell.width = Mm(122.4)

    cell_logo = table.cell(0, 0)
    cell_title = table.cell(0, 1)

    tcPr = cell_logo._tc.get_or_add_tcPr()
    tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="0" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tcMar>' % nsdecls('w'))
    tcPr.append(tcMar)

    p_logo = cell_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(0)
    p_logo.paragraph_format.line_spacing = 1
    cell_logo.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    
    try:
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Mm(57.6)) 
    except Exception as e:
        p_logo.add_run("[Logo Bulunamadı]")

    p_title = cell_title.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell_title.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    
    run_title = p_title.add_run("BİLGİ NOTU")
    run_title.font.name = 'Times New Roman'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0xC9, 0x22, 0x28) 

    cell_konu_label = table.cell(1, 0)
    cell_konu_label.paragraphs[0].add_run("Konu").font.bold = True
    
    cell_konu_val = table.cell(1, 1)
    cell_konu_val.paragraphs[0].add_run(topic)

    cell_date_label = table.cell(2, 0)
    cell_date_label.paragraphs[0].add_run("Hazırlanış Tarihi").font.bold = True
    
    cell_date_val = table.cell(2, 1)
    cell_date_val.paragraphs[0].add_run(date)

    cell_section = table.cell(3, 0)
    cell_section.merge(table.cell(3, 1)) 
    p_sec = cell_section.paragraphs[0]
    p_sec.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sec = p_sec.add_run("AÇIKLAMA")
    run_sec.font.size = Pt(12)
    run_sec.font.bold = True

    cell_content = table.cell(4, 0)
    cell_content.merge(table.cell(4, 1))
    
    table.rows[4].height = Pt(390) 
    table.rows[4].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    
    cell_content.text = "" 
    
    paragraphs = text.split('\n')
    for i, para_text in enumerate(paragraphs):
        para_text = para_text.strip()
        if not para_text:
            continue
        
        p = cell_content.add_paragraph(para_text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY  
        p.paragraph_format.first_line_indent = Mm(7) 
        
        if i < len(paragraphs) - 1:
            p.paragraph_format.space_after = Pt(12)
        else:
            p.paragraph_format.space_after = Pt(0) 

    if len(cell_content.paragraphs) > 1 and cell_content.paragraphs[0].text.strip() == "":
        p_elem = cell_content.paragraphs[0]._element
        p_elem.getparent().remove(p_elem)
        
    document.save(output_filename)

def get_kaza_kirim_raporu(licensePlate, date):
    url = f"{BASE_URL}/kirim-raporu"
    params = {
        "licensePlate": licensePlate,
        "date": date  
    }
    try:
        response = requests.get(url, params=params, timeout=15)
        if response.status_code == 200:
            data = response.json()
            return data.get("report")
        else:
            return ""
    except Exception as e:
        print("Kırım Raporu Bağlantı Hatası:", e)
        return ""

def get_konum_zamani(licensePlate):
    url = f"{BASE_URL2}/konum-zamani"
    params = {
        "licensePlate": licensePlate
    }
    try:
        response = requests.get(url, params=params, timeout=15)
        if response.status_code == 200:
            data = response.json()
            return data.get("location_time")
        else:
            return ""
    except Exception as e:
        print("Konum Zamanı Bağlantı Hatası:", e)
        return ""

@app.post("/api/haber-analiz")
async def analyze_news(request_data: HaberRequest):
    global app_data 

    
    try:
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}
        response = requests.get(request_data.url, headers=headers, timeout=15)
        response.raise_for_status()

        soup = BeautifulSoup(response.text, "html.parser")
        for tag in soup(["script", "style", "head", "svg", "noscript", "iframe"]):
            tag.decompose()
        sample_text = " ".join(soup.stripped_strings)
        
        payload = {
            "model": MODEL_NAME,
            "messages": [
                {
                    "role": "system",
                    "content": (
                        "You are an expert official incident analyst AI writing formal reports in Turkish. "
                        "Extract information and output **only** a JSON object with these exact keys:\n"
                        "{\n"
                        '  "bus_license_plates": [],\n'
                        '  "bus_count": 0,\n'
                        '  "city": "...",\n'
                        '  "district": "...",\n'
                        '  "location": "...",\n'
                        '  "accident_description": "...",\n'
                        '  "death_count": 0,\n'
                        '  "injured_count": 0,\n'
                        '  "news_date": "DD.MM.YYYY",\n'
                        '  "news_date2": "DD/MM/YYYY",\n'
                        '  "news_date3": "YYYY-MM-DD",\n'
                        '  "summary_sentences": []\n'
                        "}\n\n"
                        "CRITICAL EXTRACTION RULES:\n"
                        "1. First, identify if the text describes a crash involving ANY type of bus (otobüs, İETT otobüsü, Özel Halk Otobüsü, yolcu otobüsü, vb.). If yes, set 'bus_count' to the correct number. DO NOT rely on the presence of a license plate to count the bus. If there is a bus crash, 'bus_count' MUST be at least 1.\n"
                        "2. For each identified bus, extract its license plate. If a bus is mentioned but its license plate is missing from the text, you MUST insert an empty string (\"\") into the 'bus_license_plates' array to match the 'bus_count'.\n"
                        "3. CRITICAL LICENSE PLATE FORMAT RULE: Only when a license plate is actually present, format it with no spaces (e.g., '06ANKARA06' or '34ABC123').\n"
                        "4. If the exact accident date is missing, leave 'news_date' fields as empty strings (\"\").\n"
                        "5. CRITICAL PLACEHOLDER RULE: If the license plate or date is missing, you MUST literally write '[license_plate]' and '[news_date]' in the summary sentence. NEVER delete these placeholders.\n\n"
                        "GOLDEN TEMPLATE FOR 'summary_sentences' (Generate exactly one sentence per bus_count):\n"
                        "'Bazı basın yayın organlarında yer alan haberlerde; [news_date] tarihinde [city]nin [district] ilçesinde [license_plate] plakalı yolcu otobüsünün [location] mevkiinde [accident_description] sonucu gerçekleşen kazada [casualty_text] bilgisi yer almaktadır.'\n\n"
                        "CONDITIONAL DYNAMIC RULES:\n"
                        "1. CITY/DISTRICT MISSING: Omit '[city]nin [district] ilçesinde'. (If district is missing but city exists, use '[city]de') (If city is missing but district exists, use '[district] ilçesinde').\n"
                        "2. LOCATION MISSING: Omit '[location] mevkiinde'.\n"
                        "3. ACCIDENT DESCRIPTION MISSING: Omit '[accident_description] sonucu gerçekleşen kazada' and connect seamlessly.\n"
                        "4. CASUALTY TEXT RULES (CRITICAL):\n"
                        "   - If injured > 0 AND dead > 0: '{injured_count} vatandaşımızın yaralandığı, {death_count} vatandaşımızın ise vefat ettiği'\n"
                        "   - If injured > 0 AND dead == 0: '{injured_count} vatandaşımızın yaralandığı'\n"
                        "   - If dead > 0 AND injured == 0: '{death_count} vatandaşımızın vefat ettiği'\n"
                        "   - If dead == 0 AND injured == 0: 'can kaybı ya da yaralanmanın olmadığı'\n"
                        "   NEVER write '0 ölü' or '0 yaralı'. If a count is zero, completely omit that specific part.\n"
                        "5. ACCIDENT DESCRIPTION GRAMMAR (CRITICAL): Do NOT use past tense verbs. Convert the final action into an 'isim-fiil' (verbal noun) ending with '-ması' or '-mesi' (e.g., 'duvara çarpması', 'devrilmesi').\n"
                        "Ensure perfect Turkish grammar."
                    )
                },
                {
                    "role": "user",
                    "content": f"Analyze this text and extract details into JSON:\n\n{sample_text[:10000]}"
                }
            ],
            "temperature": 0.1,
            "top_p": 0.1,
            "max_tokens": 4096,
            "chat_template_kwargs": {"enable_thinking": False}
        }

        payload_data = json.dumps(payload).encode('utf-8')
        req = urllib.request.Request(URL, data=payload_data, headers={'Content-Type': 'application/json'})

        with urllib.request.urlopen(req) as response:
            result = json.loads(response.read().decode('utf-8'))
            
            message_obj = result.get('choices', [{}])[0].get('message', {})
            raw_content = message_obj.get('content') if message_obj else None
            
            if not raw_content:
                raise ValueError("Model boş yanıt döndürdü.")
            
            clean_json = raw_content.replace("```json", "").replace("```", "").strip()
            parsed_data = json.loads(clean_json)

            app_data["bus_count"] = parsed_data.get("bus_count", 0)
            raw_plates = parsed_data.get("bus_license_plates", [])
            app_data["bus_license_plates"] = [p.replace(" ", "").upper() for p in raw_plates if p]
            app_data["news_date"] = parsed_data.get("news_date", "")
            app_data["news_date2"] = parsed_data.get("news_date2", "")
            app_data["news_date3"] = parsed_data.get("news_date3", "")
            app_data["summary_sentences"] = parsed_data.get("summary_sentences", [])


            return {
                "busPlates": app_data["bus_license_plates"],
                "newsDate": app_data["news_date3"],
                "busCount": app_data["bus_count"]
            }

    except Exception as e:
        print(f"[CHECKER HATA] Analiz Sırasında Hata Oluştu: {str(e)}")
        raise HTTPException(status_code=500, detail="Haber analiz edilirken bir hata oluştu.")


@app.post("/api/rapor-olustur")
def create_reports(request_data: RaporRequest):
    global app_data
    generated_files = []
    
    try:
        image_path = os.path.join(os.path.dirname(__file__), "logo.png")
        if not os.path.exists(image_path):
            raise HTTPException(status_code=500, detail="Kurumsal logo (logo.png) bulunamadığı için rapor oluşturulamadı.")

        for i, vehicle in enumerate(request_data.buses):
            current_plate = vehicle.plate.strip().upper()
            licence_plate_formatted = current_plate.replace(" ", "")
            
            news_date3 = vehicle.accident_date if vehicle.accident_date else app_data.get("news_date3", "")
            news_date = app_data.get("news_date", "")
            news_date2 = app_data.get("news_date2", "")
            
            if news_date3 and "-" in news_date3:
                year, month, day = news_date3.split("-")
                news_date = f"{day}.{month}.{year}"
                news_date2 = f"{day}/{month}/{year}"
            

            if i < len(app_data["summary_sentences"]):
                current_summary = app_data["summary_sentences"][i]
            else:
                current_summary = "Bazı basın yayın organlarında yer alan haberlerde; [news_date] tarihinde [license_plate] plakalı yolcu otobüsünün kazaya karıştığı bilgisi yer almaktadır."

            old_plate = app_data["bus_license_plates"][i] if i < len(app_data["bus_license_plates"]) else ""
            old_date = app_data["news_date"]

            if old_plate and old_plate in current_summary:
                current_summary = current_summary.replace(old_plate, current_plate)
            elif "[license_plate]" in current_summary:
                current_summary = current_summary.replace("[license_plate]", current_plate)
                
            if old_date and old_date in current_summary:
                current_summary = current_summary.replace(old_date, news_date)
            elif "[news_date]" in current_summary:
                current_summary = current_summary.replace("[news_date]", news_date)

            first_text = "\n" + current_summary + "\n\n"

            try:
                second_text = get_kaza_kirim_raporu(licence_plate_formatted, news_date3)
                location_text = get_konum_zamani(licence_plate_formatted)
            except ConnectionError:
                print(f"[CHECKER HATA] Veritabanı bağlantısı kurulamadı! 503 döndürülüyor.")
                raise HTTPException(status_code=503, detail="Raporlar oluşturulamadı")
            
            if second_text or location_text:
                if second_text:
                    second_text = second_text.replace("\r\n", "\n").replace("\r", "\n")
                else:
                    second_text = ""
                
                if location_text:
                    location_text = location_text.replace("\r\n", "\n").replace("\r", "\n")
                else:
                    location_text = ""
                
                raw_final_text = (second_text + "\n" + location_text).strip()
            else:
                raw_final_text = "Herhangi bir kayda ulaşılamamıştır."

            formatter_payload = {
                "model": MODEL_NAME,
                "messages": [
                    {
                        "role": "system",
                        "content": (
                            "You are an expert official incident report editor for a Turkish Government Ministry. Your task is to take a raw database output string regarding a vehicle's registration, inspection, insurance, U-ETDS trip data, and ATS location tracking data, and rewrite it STRICTLY into the formal official template provided below.\n\n"
                            "=========================================\n"
                            "TARGET FORMAL TEMPLATE STRUCTURE (NORMAL SCENARIO)\n"
                            "=========================================\n"
                            "Söz konusu kazaya karışan [PLAKA] plakalı taşıt hakkında, Bakanlığımız kayıtlarında yapılan inceleme neticesinde;\n"
                            "➢ Bahse konu taşıtın, [Firma Adı] firması adına düzenlenen [Belge Türü] yetki belgesinde [Özmal / Sözleşmeli (Tescilli Firma Adı adına tescilli)] olarak kayıtlı bulunduğu,\n"
                            "➢ Anılan taşıtın, [Model Yılı] model otobüs cinsi olduğu ve [Koltuk Sayısı] koltukluk taşıma kapasitesine sahip olduğu,\n"
                            "➢ Taşıtın, [Muayene Tarihi] tarihine kadar geçerli araç muayenesinin bulunduğu,\n"
                            "➢ Taşıtın, [Trafik Sigortası Durumu] ve [Koltuk Sigortası Durumu] bulunduğu,\n"
                            "tespit edilmiştir.\n\n"
                            "Kazayla ilgili olarak Bakanlığımız U-ETDS sisteminde yapılan incelemede;\n"
                            "➢ Anılan taşıtla gerçekleştirilen seferin [Tarih] tarihinde saat [Saat]'da [Kalkış Yeri] başladığı ve [Varış Yeri] sona ereceğinin bildirildiği,\n"
                            "➢ Seferle ilgili en son [Tarih] tarihinde saat [Saat]'de U-ETDS sistemine veri iletildiği,\n"
                            "➢ Bahse konu sefer için U-ETDS sistemine toplam [Yolcu Sayısı] yolcunun iletildiği (ara duraklardaki indirme ve bindirmeler dahil),\n"
                            "➢ Anılan seferde, sürücü [Sürücü1 Ad Soyad]'ın ([SRC Belgesi]) otobüse biniş yerinin [Yer] olarak iletildiği, sürücü [Sürücü2 Ad Soyad]'un ([SRC Belgesi]) otobüse biniş yerinin [Yer] olarak iletildiği ve yardımcı personel [Personel Ad Soyad]'ın iletildiği,\n"
                            "➢ Bahse konu araç için en son [Tarih Saat] tarihinde bakanlığımız sistemine konum verisi iletildiği (VEYA ➢ Bahse konu araç için bakanlığımız sistemine konum verisi iletilmediği,)\n"
                            "tespit edilmiştir.\n"
                            "Arz ederim.\n\n"
                            "=========================================\n"
                            "CRITICAL TRANSFORMATION RULES & SCENARIOS\n"
                            "=========================================\n"
                            "1. FIRST SENTENCE (MANDATORY): You MUST ALWAYS start the report with this exact phrase: \n"
                            f'   "Söz konusu kazaya karışan {current_plate} plakalı taşıt hakkında, Bakanlığımız kayıtlarında yapılan inceleme neticesinde;"\n'
                            "2. BULLET POINTS: Use ONLY the `➢` character for bullet points. Do NOT use `*` or `-` for list markers. Convert all `*` from the raw text to `➢`.\n"
                            "3. SCENARIO A: COMPLETE VEHICLE RECORD NOT FOUND (e.g., \"Herhangi bir kayda ulaşılamamıştır\", \"Taşıt yetki belgesi kaydı bulunmamaktadır\"): \n"
                            "   - Do NOT generate the registration/insurance bullets or the U-ETDS section. \n"
                            "   - Format it strictly as:\n"
                            "     ➢ Herhangi bir kayda ulaşılamamıştır,\n"
                            "     tespit edilmiştir.\n"
                            "     Arz ederim.\n"
                            "5. SCENARIO B & C: VEHICLE RECORD EXISTS (Normal Transformation Rules):\n"
                            "   - CERTIFICATE TYPE & OWNERSHIP: Extract the exact certificate type from the raw text. If it says \"Özmal\", write \"Özmal\". If it says \"Sözleşmeli\" with a company, write it strictly as \"Sözleşmeli ([Firma Adı] adına tescilli)\".\n"
                            "   - MISSING/DIFFERENT DATA (INSURANCE & INSPECTION): Handle Traffic Insurance, Compulsory Seat Accident Insurance, and Inspection independently based on the raw text.\n"
                            "   - U-ETDS SECTION & MISSING TRIPS: The second section MUST start with: \"Kazayla ilgili olarak Bakanlığımız U-ETDS sisteminde yapılan incelemede;\"\n"
                            "     * If the raw text states that there is no trip (e.g., \"UETDS Sefer bildirimi bulunmamaktadır\"), format it strictly and cleanly as: `➢ UETDS Sefer bildirimi bulunmamaktadır,`\n"
                            "     * If trip details exist, format them using all the detailed bullets.\n"
                            "   - PERSONNEL MERGING: Merge drivers and personnel into a single, fluent bullet point as shown in the template.\n"
                            "   - ATS / LOCATION TRACKING POSITION: If raw text contains location/ATS data, it MUST be placed as a bullet point (`➢`) *before* \"tespit edilmiştir.\" and *before* \"Arz ederim.\".\n"
                            "     * Format if data exists: `➢ Bahse konu araç en son [Tarih Saat] tarihinde konum bilgisi vermiştir.`\n"
                            "     * Format if data is missing: `➢ Bahse konu araca ait konum bilgisi bulunamadığı,`\n"
                            "   - ENDING: The text MUST end strictly with:\n"
                            "     tespit edilmiştir.\n"
                            "     Arz ederim.\n\n"
                            "Ensure perfect Turkish administrative grammar and output ONLY the final formatted text."
                        )
                    },
                    {
                        "role": "user",
                        "content": f"Rewrite this raw database output into the formal official template:\n\n{raw_final_text}"
                    }
                ],
                "temperature": 0.3,
                "top_p": 0.1,
                "max_tokens": 4096,
                "chat_template_kwargs": {"enable_thinking": False}
            }

            formatter_payload_data = json.dumps(formatter_payload).encode('utf-8')
            formatter_req = urllib.request.Request(URL, data=formatter_payload_data, headers={'Content-Type': 'application/json'})
            
            with urllib.request.urlopen(formatter_req) as formatter_response:
                formatter_result = json.loads(formatter_response.read().decode('utf-8'))
                
                f_message_obj = formatter_result.get('choices', [{}])[0].get('message', {})
                final_content = f_message_obj.get('content') if f_message_obj else None
                
                if not final_content:
                    print(f"   [CHECKER HATA] {current_plate} için LLM boş içerik döndürdü!")
                    continue
                    
                final_text = final_content.strip()

            pdf_filename = f"{licence_plate_formatted}_KazaRaporu.pdf"
            docx_filename = f"{licence_plate_formatted}_KazaRaporu.docx"
            
            pdf_path = os.path.join("reports", pdf_filename)
            docx_path = os.path.join("reports", docx_filename)
            
            final_docx_text = first_text + final_text
            
            generate_pdf_report(pdf_path, current_plate, news_date2, final_text, image_path)

            generate_formal_report_docx(docx_path, current_plate, news_date2, final_docx_text, image_path)

            generated_files.append({
                "pdfUrl": f"http://localhost:8000/reports/{pdf_filename}",
                "docxUrl": f"http://localhost:8000/reports/{docx_filename}",
                "fileName": f"{licence_plate_formatted}_KazaRaporu"
            })

        return {"raporlar": generated_files}

    except HTTPException as he:
        raise he
    except Exception as e:
        print(f"[CHECKER HATA] Rapor Oluşturulurken Kritik Hata: {str(e)}")
        raise HTTPException(status_code=500, detail="Raporlar oluşturulurken bir hata meydana geldi.")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)